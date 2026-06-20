"""白箱 SAST、標準對應、案件報告與跨次追蹤的測試。"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from pentest import docx_report, sast
from pentest.checks.base import Finding, Severity
from pentest.engagement import Engagement
from pentest.standards import enrich, kind_from_check_id, owasp_coverage

SAMPLE = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      "examples", "vulnerable_app")


def test_sast_finds_volvo_style_vulns():
    findings = sast.scan_path(SAMPLE, target_name="demo")
    kinds = {f.check_id.split("-")[1] for f in findings if f.severity != Severity.INFO}
    for expected in ("sqli", "xss", "secret", "dep", "cmdi"):
        assert expected in kinds, f"未偵測到 {expected};實得 {kinds}"


def test_sast_findings_are_enriched_with_standards():
    findings = sast.scan_path(SAMPLE, target_name="demo")
    sqli = [f for f in findings if "sqli" in f.check_id and f.severity == Severity.CRITICAL]
    assert sqli
    f = sqli[0]
    assert f.owasp.startswith("A03")
    assert f.cwe == "CWE-89"
    assert f.cvss and f.cvss >= 9.0
    assert ":" in f.location          # 檔案:行號
    assert f.stable_id                # 有穩定 ID


def test_kind_from_check_id():
    assert kind_from_check_id("active-sqli-id") == "sqli"
    assert kind_from_check_id("header-missing-csp") == "header"
    assert kind_from_check_id("active-openredirect-next") == "openredirect"


def test_owasp_coverage_matrix():
    findings = sast.scan_path(SAMPLE, target_name="demo")
    cov = owasp_coverage(findings)
    assert cov["A03"]      # Injection 應被涵蓋
    assert cov["A06"]      # Vulnerable Components(xlsx CVE)


def test_enrich_sets_cvss_band():
    f = enrich(Finding("x", "t", Severity.HIGH, "d", "r"), "xss")
    assert f.cvss == 7.5
    assert f.cwe == "CWE-79"


def test_report_generation(tmp_path):
    findings = sast.scan_path(SAMPLE, target_name="demo")
    eng = Engagement(target="demo-app", name="Demo", methodology="white-box",
                     test_type="SAST", tester="tester")
    out = str(tmp_path / "report.docx")
    docx_report.generate(eng, findings, out)
    assert os.path.exists(out)
    doc = Document(out)
    heads = [p.text for p in doc.paragraphs]
    joined = "\n".join(heads)
    assert "滲透測試報告書" in joined
    assert "OWASP Top 10" in joined
    assert "修補路線圖" in joined
    assert "測試整合性與清理聲明" in joined


def test_report_shows_authenticated_coverage_section(tmp_path):
    """灰箱掃描:報告應有『測試覆蓋與已認證攻擊面』區塊,攤開前端/後端/登入後內部測了什麼。"""
    import json
    cov = {"authenticated": True, "pages": 6, "forms": 1, "points": 3, "api_count": 2,
           "api_list": ["/api/daily", "/api/users/me"], "active": True,
           "browser_used": True, "authz_tested": 2, "bac_hits": 1}
    findings = [
        enrich(Finding("authz-bac-api-daily", "缺少授權驗證:未登入即可存取內部 API — /api/daily",
                       Severity.HIGH, "未帶 session 仍回資料。", "強制授權。"), "accesscontrol"),
        Finding("scan-coverage", "測試覆蓋摘要", Severity.INFO,
                json.dumps(cov, ensure_ascii=False), "—"),
    ]
    eng = Engagement(target="https://x/", name="t", methodology="grey-box",
                     test_type="DAST", tester="s")
    out = str(tmp_path / "r.docx")
    docx_report.generate(eng, findings, out)
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    assert "測試覆蓋與已認證攻擊面" in joined
    assert "灰箱" in joined and "/api/daily" in joined
    assert "缺少授權驗證" in joined                       # 登入後內部結論有呈現
    # BAC finding 自身也在 A01 並會出現在 Findings 細節
    assert any(f.owasp.startswith("A01") for f in findings if f.check_id.startswith("authz-bac-"))


def test_report_honest_when_no_internal_api(tmp_path):
    """已登入但沒探到內部 API → 報告誠實說明原因與補法,而非假裝測過。"""
    import json
    cov = {"authenticated": True, "pages": 2, "forms": 1, "points": 1, "api_count": 0,
           "api_list": [], "active": True, "browser_used": False,
           "authz_tested": 0, "bac_hits": 0}
    findings = [Finding("scan-coverage", "測試覆蓋摘要", Severity.INFO,
                        json.dumps(cov, ensure_ascii=False), "—")]
    eng = Engagement(target="https://x/", name="t", methodology="grey-box",
                     test_type="DAST", tester="s")
    out = str(tmp_path / "r2.docx")
    docx_report.generate(eng, findings, out)
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    assert "未能探出可測的內部 API 端點" in joined
    assert "瀏覽器動態爬取" in joined          # 給出補法


def test_new_recurring_fixed_status(tmp_path):
    # 上次稽核有 a,b;這次有 b,c → b=recurring, c=new, a=fixed
    findings = [
        Finding("sast-sqli-x", "T-B", Severity.HIGH, "d", "r", stable_id="t::sqli::b"),
        Finding("sast-xss-x", "T-C", Severity.MEDIUM, "d", "r", stable_id="t::xss::c"),
    ]
    eng = Engagement(target="t", name="T", methodology="white-box", test_type="SAST")
    out = str(tmp_path / "r.docx")
    docx_report.generate(eng, findings, out,
                         previous_stable_ids={"t::sqli::b", "t::sqli::a"})
    status = {f.stable_id: f.status for f in findings}
    assert status["t::sqli::b"] == "recurring"
    assert status["t::xss::c"] == "new"
    # 報告應列出已修復的 a
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    assert "t::sqli::a" in joined


def test_cleanup_attestation_no_artifacts():
    eng = Engagement(target="t", name="T")
    report = eng.cleanup_report()
    assert report["persistent_changes"] is False
    assert "反鑑識" in report["note"]


def test_build_context_and_print_report(tmp_path):
    """build_context 彙整資料 + 列印報告渲染器(與 .docx 同源)。"""
    import json
    from pentest import docx_report
    cov = {"authenticated": True, "pages": 5, "forms": 1, "points": 2, "api_count": 1,
           "api_list": ["/api/daily"], "active": True, "browser_used": True,
           "authz_tested": 1, "bac_hits": 1}
    findings = [
        enrich(Finding("authz-bac-api-daily", "缺少授權驗證 — /api/daily", Severity.HIGH,
                       "未帶 session 仍回資料。", "強制授權。"), "accesscontrol"),
        Finding("scan-coverage", "測試覆蓋摘要", Severity.INFO,
                json.dumps(cov, ensure_ascii=False), "—"),
    ]
    eng = Engagement(target="https://x/", name="t", methodology="grey-box",
                     test_type="DAST", tester="s")
    ctx = docx_report.build_context(eng, findings)
    assert ctx["method_label"] == "灰箱"
    assert ctx["counts"][Severity.HIGH] == 1
    assert ctx["coverage"]["api_count"] == 1
    assert any(f.check_id.startswith("authz-bac-") for f in ctx["coverage_bac"])
    assert "flow" in ctx and "owasp" in ctx and ctx["roadmap"]


def test_report_checklist_and_tested_clean_matrix(tmp_path):
    """測試項目清單 +『已測試,未發現』矩陣:讓 SQLi/上傳/後端等『測了沒漏洞』也看得見。"""
    import json
    cov = {"authenticated": True, "pages": 7, "forms": 2, "points": 2, "api_count": 3,
           "api_list": ["/api/users"], "active": True, "browser_used": True,
           "authz_tested": 3, "bac_hits": 0,
           "owasp_tested": ["A01", "A03", "A05"],
           "checklist": [
               {"id": "active", "label": "主動弱點測試(注入 / XSS / 轉址)", "findings": 0, "ran": True},
               {"id": "upload", "label": "檔案上傳攻擊面", "findings": 0, "ran": True},
               {"id": "headers", "label": "安全回應標頭", "findings": 1, "ran": True}]}
    findings = [
        enrich(Finding("header-missing-csp", "缺少安全標頭:CSP", Severity.MEDIUM, "d", "r"), "header"),
        Finding("scan-coverage", "測試覆蓋摘要", Severity.INFO, json.dumps(cov, ensure_ascii=False), "—"),
    ]
    eng = Engagement(target="https://x/", name="t", methodology="grey-box",
                     test_type="DAST", tester="s")
    ctx = docx_report.build_context(eng, findings)
    assert "A03" in ctx["owasp_tested"] and len(ctx["checklist"]) == 3
    out = str(tmp_path / "r.docx")
    docx_report.generate(eng, findings, out)
    joined = "\n".join(p.text for p in Document(out).paragraphs)
    assert "測試項目與覆蓋" in joined
    # 矩陣應對「已測試但無 finding」標示「已測試,未發現」,對未測類別標「未涵蓋」
    cells = [c.text for tb in Document(out).tables for r in tb.rows for c in r.cells]
    assert any("已測試,未發現" in x for x in cells)
    assert any("未涵蓋" in x for x in cells)
